from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import pdfplumber
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
MAPPING_FILE = BASE_DIR / "mapeamento_enderecos.csv"
RECEIVING_HEADERS = [
    "TRECHO", "PONTO ORIGEM", "PONTO DESTINO", "MANIFESTO", "VOLUMES",
    "DATA DE TRASFERENCIA", "VOL. RECEBIDOS", "VOLUMES NÃO RECEBIDO",
    "CTE'S FALTANTE", "STATUS CTE'S FALTANTE", "AGING",
    "STATUS DE COBRANÇA", "CANAL", "DT DA COBRANÇA",
]
SHIPPING_HEADERS = [
    "PONTO ORIGEM", "PONTO DESTINO", "DT DE TRASFERENCIA", "MANIFESTO",
    "VOLUMES", "VOL. NÃO RECEBIDO", "CTE'S NÃO RECEBIDO", "STATUS",
    "DT DO TRANSFERIDO", "TRECHO", "AGING", "VEICULO", "SM",
    "HORA DE CHEGADA", "STATUS DE COBRANÇA", "CANAL", "DT DA COBRANÇA",
]
TARGET_ADDRESS = "RUA COMANDANTE VERGUEIRO DA CRUZ"
DESTINATION_FIXED = "FL RIO DEVOLUÇÃO"
MONTHS = {
    "jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
    "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12,
}


@dataclass
class Manifest:
    filename: str
    number: str
    transfer_date: date | None
    volumes: int
    ctes: list[str]
    origin_text: str
    destination_text: str
    pages: int


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def clean_duplicated_line(line: str) -> str:
    letters = [c for c in line if c.isalpha()]
    if not letters:
        return line
    pairs = sum(1 for a, b in zip(line, line[1:]) if a.isalpha() and a == b)
    if pairs / max(1, len(letters)) < 0.30:
        return line
    return re.sub(r"(.)\1+", r"\1", line)


def cleaned_text(page: pdfplumber.page.Page) -> str:
    raw = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
    return "\n".join(clean_duplicated_line(line) for line in raw.splitlines())


def grouped_word_lines(page: pdfplumber.page.Page) -> list[tuple[float, list[dict[str, Any]]]]:
    words = page.extract_words(x_tolerance=2, y_tolerance=3, keep_blank_chars=False)
    words.sort(key=lambda word: (word["top"], word["x0"]))
    groups: list[tuple[float, list[dict[str, Any]]]] = []
    for word in words:
        if groups and abs(groups[-1][0] - word["top"]) <= 3:
            groups[-1][1].append(word)
        else:
            groups.append((word["top"], [word]))
    return groups


def party_blocks(page: pdfplumber.page.Page) -> tuple[str, str]:
    """Read the two address blocks by their left/right page positions."""
    lines = grouped_word_lines(page)
    header_y = None
    table_y = None
    for y, words in lines:
        line = normalize(clean_duplicated_line(" ".join(w["text"] for w in words)))
        if header_y is None and "EMITENTE" in line and "DESTINATARIO" in line:
            header_y = y
        if header_y is not None and y > header_y and "REMESSA" in line:
            table_y = y
            break
    if header_y is None:
        # Some PDF text layers place the two labels in separate y-groups.
        for y, words in lines:
            line = normalize(clean_duplicated_line(" ".join(w["text"] for w in words)))
            if "EMITENTE" in line:
                header_y = y
                break
    if header_y is None:
        return "", ""
    if table_y is None:
        table_y = page.height * 0.42
    midpoint = page.width / 2
    sides: list[list[str]] = [[], []]
    for y, words in lines:
        if y <= header_y + 5 or y >= table_y:
            continue
        for side, selected in enumerate(([w for w in words if w["x0"] < midpoint],
                                         [w for w in words if w["x0"] >= midpoint])):
            if selected:
                sides[side].append(" ".join(w["text"] for w in selected))
    return " ".join(sides[0]), " ".join(sides[1])


def parse_manifest(uploaded_file: Any) -> Manifest:
    with pdfplumber.open(uploaded_file) as pdf:
        if not pdf.pages:
            raise ValueError("O PDF não contém páginas.")
        all_text = [cleaned_text(page) for page in pdf.pages]
        first = all_text[0]
        first_page = pdf.pages[0]
        origin, destination = party_blocks(first_page)

        # The first 14-digit ID in the header is the manifesto number.
        before_parties = re.split(r"EMITENTE", first, maxsplit=1, flags=re.IGNORECASE)[0]
        number_match = re.search(r"\b\d{14}\b", before_parties)
        if not number_match:
            top_words = [w["text"] for w in first_page.extract_words()
                         if w["top"] < first_page.height * 0.20]
            number_match = re.search(r"\b\d{14}\b", " ".join(top_words))
        if not number_match:
            raise ValueError("Não encontrei o número de 14 dígitos do manifesto no cabeçalho.")
        number = number_match.group(0)

        date_match = re.search(
            r"\b\d{1,2}\s+(jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)\s+\d{4}\b",
            normalize_for_date(first),
        )
        transfer_date = None
        if date_match:
            day, month, year = date_match.group(0).split()
            transfer_date = date(int(year), MONTHS[month], int(day))

        ctes: list[str] = []
        volume_total = 0
        for page, text in zip(pdf.pages, all_text):
            page_total_match = re.search(r"VOLUMES?\s*:?\s*(\d+)\b", normalize(text))
            if page_total_match:
                volume_total += int(page_total_match.group(1))
            page_rows = 0
            for line in text.splitlines():
                row_match = re.match(r"\s*(\d{14})\s*/\s*\d{14}\b(.*)$", line)
                if not row_match:
                    continue
                cte = row_match.group(1)
                if cte not in ctes:
                    ctes.append(cte)
                page_rows += 1
                # Fallback if the page has no printed total: the denominator in VOL (e.g. 1/2).
                if not page_total_match:
                    vol_match = re.search(r"\b(\d+)\s*/\s*(\d+)\b", row_match.group(2))
                    volume_total += int(vol_match.group(2)) if vol_match else 1
            if page_rows == 0 and not page_total_match:
                raise ValueError(f"Não encontrei linhas de CT-e na página {page.page_number}.")

    return Manifest(
        filename=uploaded_file.name,
        number=number,
        transfer_date=transfer_date,
        volumes=volume_total,
        ctes=ctes,
        origin_text=origin,
        destination_text=destination,
        pages=len(all_text),
    )


def normalize_for_date(text: str) -> str:
    return normalize(text).lower()


def load_mapping() -> list[dict[str, str]]:
    with MAPPING_FILE.open(encoding="utf-8-sig", newline="") as f:
        return [
            {
                "Endereço que pode aparecer no manifesto": row[0],
                "Rótulo que deve sair na planilha": row[1],
                "Trecho": row[2],
            }
            for row in list(csv.reader(f))[1:]
            if len(row) >= 2 and row[0].strip()
        ]


def mapping_aliases(mapping_rows: list[dict[str, Any]]) -> list[tuple[str, str, str]]:
    aliases = []
    for row in mapping_rows:
        address = str(row.get("Endereço que pode aparecer no manifesto", "") or "")
        label = str(row.get("Rótulo que deve sair na planilha", "") or "")
        section = str(row.get("Trecho", "") or "")
        for alias in re.split(r"\s+OU\s+|\s*;\s*|\s*\|\s*", address, flags=re.IGNORECASE):
            key = normalize(alias)
            if key and label.strip():
                aliases.append((key, label.strip(), section.strip()))
    return sorted(aliases, key=lambda item: len(item[0]), reverse=True)


def match_point(text: str, aliases: list[tuple[str, str, str]]) -> tuple[str, str] | None:
    normalized = normalize(text)
    for key, label, section in aliases:
        if key and re.search(rf"\b{re.escape(key)}\b", normalized):
            return label, section
    return None


def make_output_row(manifest: Manifest, kind: str, mapping: list[tuple[str, str, str]]) -> tuple[dict[str, Any] | None, list[str]]:
    warnings: list[str] = []
    origin_target = TARGET_ADDRESS in normalize(manifest.origin_text)
    destination_target = TARGET_ADDRESS in normalize(manifest.destination_text)
    if origin_target == destination_target:
        reason = "aparece em ORIGEM e DESTINO" if origin_target else "não aparece nem em ORIGEM nem em DESTINO"
        return None, [f"{manifest.filename}: não foi possível classificar; o endereço de referência {reason}."]

    if destination_target:
        kind = "recebendo"
        other_text = manifest.origin_text
        point = match_point(other_text, mapping)
        if point is None:
            warnings.append(f"{manifest.filename}: endereço de ORIGEM sem correspondência; PONTO ORIGEM ficará em branco.")
        receiving_row = {
            "TRECHO": "", "PONTO ORIGEM": point[0] if point else "",
            "PONTO DESTINO": DESTINATION_FIXED, "MANIFESTO": manifest.number,
            "VOLUMES": manifest.volumes, "DATA DE TRASFERENCIA": manifest.transfer_date,
            "VOL. RECEBIDOS": "", "VOLUMES NÃO RECEBIDO": len(manifest.ctes),
            "CTE'S FALTANTE": "; ".join(manifest.ctes),
            "STATUS CTE'S FALTANTE": "PENDENTE DE RECEBIMENTO", "AGING": "",
            "STATUS DE COBRANÇA": "", "CANAL": "", "DT DA COBRANÇA": "",
        }
        return receiving_row, warnings

    kind = "expedindo"
    point = match_point(manifest.destination_text, mapping)
    if point is None:
        warnings.append(f"{manifest.filename}: endereço de DESTINO sem correspondência; PONTO DESTINO ficará em branco.")
    shipping_row = {
        "PONTO ORIGEM": DESTINATION_FIXED, "PONTO DESTINO": point[0] if point else "",
        "DT DE TRASFERENCIA": manifest.transfer_date, "MANIFESTO": manifest.number,
        "VOLUMES": manifest.volumes, "VOL. NÃO RECEBIDO": len(manifest.ctes),
        "CTE'S NÃO RECEBIDO": "; ".join(manifest.ctes), "STATUS": "NO PISO DA FL",
        "DT DO TRANSFERIDO": "", "TRECHO": "", "AGING": "", "VEICULO": "",
        "SM": "", "HORA DE CHEGADA": "", "STATUS DE COBRANÇA": "", "CANAL": "",
        "DT DA COBRANÇA": "",
    }
    return shipping_row, warnings


def build_xlsx(receiving: list[dict[str, Any]], shipping: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in (
        ("BASE-RECEBENDO", RECEIVING_HEADERS, receiving),
        ("BASE-EXPEDINDO", SHIPPING_HEADERS, shipping),
    ):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append([f'=IF(F{sheet.max_row + 1}="","",IF(TODAY()-F{sheet.max_row + 1}>=6,"D5+","D"&MAX(0,TODAY()-F{sheet.max_row + 1})))' if title == "BASE-RECEBENDO" and header == "AGING" else row.get(header, "") for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 32
        for col_idx, header in enumerate(headers, start=1):
            width = max(14, min(32, len(header) + 3))
            sheet.column_dimensions[get_column_letter(col_idx)].width = width
            if "DATA" in header:
                for row_idx in range(2, sheet.max_row + 1):
                    sheet.cell(row_idx, col_idx).number_format = "dd/mm/yyyy"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def dataframe_tsv(rows: list[dict[str, Any]], headers: list[str], include_header: bool = True) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter="\t", lineterminator="\n")
    if include_header: writer.writerow(headers)
    for row in rows:
        writer.writerow(['=SE(INDIRETO("RC[-5]";FALSO)="";"";SE(HOJE()-INDIRETO("RC[-5]";FALSO)>=6;"D5+";"D"&MÁXIMO(0;HOJE()-INDIRETO("RC[-5]";FALSO))))' if header == "AGING" and "CTE'S FALTANTE" in headers else row.get(header, "") for header in headers])
    return stream.getvalue().encode("utf-8-sig")


st.set_page_config(page_title="Manifestos Jadlog", page_icon="📦", layout="wide")
st.title("Preparar planilhas de manifestos")
st.write(
    "Envie um ou vários manifestos PDF. Cada arquivo será tratado separadamente; "
    "as páginas de um mesmo PDF serão reunidas em uma única linha."
)

if False: st.info(
    "Manifesto com Rua Comandante Vergueiro da Cruz em DESTINO → BASE-RECEBENDO. "
    "Em ORIGEM → BASE-EXPEDINDO. O CT-e é o primeiro número de 14 dígitos da coluna REMESSA."
)

uploads = st.file_uploader(
    "Selecione os manifestos PDF", type=["pdf"], accept_multiple_files=True
)

with st.expander("Conferir ou ajustar a tabela de endereços", expanded=False):
    st.caption("A tabela vem da aba Planilha2 da base enviada e inclui as correções combinadas.")
    mapping_editor = st.data_editor(
        load_mapping(),
        num_rows="dynamic",
        width="stretch",
        hide_index=True,
        column_config={
            "Endereço que pode aparecer no manifesto": st.column_config.TextColumn(required=True),
            "Rótulo que deve sair na planilha": st.column_config.TextColumn(required=True),
            "Trecho": st.column_config.TextColumn(),
        },
        key="mapping_editor",
    )

if st.button("Ler manifestos", type="primary", disabled=not uploads):
    receiving_rows: list[dict[str, Any]] = []
    shipping_rows: list[dict[str, Any]] = []
    messages: list[str] = []
    aliases = mapping_aliases(mapping_editor.to_dict(orient="records") if hasattr(mapping_editor, "to_dict") else mapping_editor)
    summaries: list[dict[str, Any]] = []
    for pdf_file in uploads:
        try:
            manifest = parse_manifest(pdf_file)
            row, warnings = make_output_row(manifest, "", aliases)
            messages.extend(warnings)
            if row is None:
                continue
            if row.get("STATUS CTE'S FALTANTE"):
                receiving_rows.append(row)
                tipo = "BASE-RECEBENDO"
                destination = row["PONTO DESTINO"]
            else:
                shipping_rows.append(row)
                tipo = "BASE-EXPEDINDO"
                destination = row["PONTO DESTINO"]
            summaries.append({
                "Arquivo": manifest.filename, "Manifesto": manifest.number,
                "Páginas": manifest.pages, "CT-e": len(manifest.ctes),
                "Volumes": manifest.volumes, "Base": tipo,
                "Ponto destino": destination,
            })
        except Exception as exc:
            messages.append(f"{pdf_file.name}: não foi possível ler o arquivo ({exc}).")

    st.session_state["manifest_receiving"] = receiving_rows
    st.session_state["manifest_shipping"] = shipping_rows
    st.session_state["manifest_messages"] = messages
    st.session_state["manifest_summaries"] = summaries

if "manifest_summaries" in st.session_state:
    summaries = st.session_state["manifest_summaries"]
    messages = st.session_state["manifest_messages"]
    if summaries:
        st.subheader("Manifestos processados")
        st.dataframe(summaries, width="stretch", hide_index=True)
        receiving_rows = st.session_state["manifest_receiving"]
        shipping_rows = st.session_state["manifest_shipping"]
        if receiving_rows is not None:
            st.subheader("Linhas para BASE-RECEBENDO")
            receiving_rows = st.data_editor(receiving_rows if receiving_rows else {header: [] for header in RECEIVING_HEADERS}, num_rows="dynamic", width="stretch", hide_index=True, key=f"receiving_results_editor_{hash(str(summaries))}"); receiving_rows = receiving_rows.to_dict(orient="records") if hasattr(receiving_rows, "to_dict") else ([dict(zip(receiving_rows, values)) for values in zip(*receiving_rows.values())] if isinstance(receiving_rows, dict) else receiving_rows); st.session_state["manifest_receiving"] = receiving_rows
        if shipping_rows is not None:
            st.subheader("Linhas para BASE-EXPEDINDO")
            shipping_rows = st.data_editor(shipping_rows if shipping_rows else {header: [] for header in SHIPPING_HEADERS}, num_rows="dynamic", width="stretch", hide_index=True, key=f"shipping_results_editor_{hash(str(summaries))}"); shipping_rows = shipping_rows.to_dict(orient="records") if hasattr(shipping_rows, "to_dict") else ([dict(zip(shipping_rows, values)) for values in zip(*shipping_rows.values())] if isinstance(shipping_rows, dict) else shipping_rows); st.session_state["manifest_shipping"] = shipping_rows
        xlsx = build_xlsx(receiving_rows, shipping_rows)
        st.download_button(
            "Baixar planilha pronta (.xlsx)", xlsx,
            file_name="manifestos_processados.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )
        left, right = st.columns(2)
        with left:
            st.download_button(
                "Baixar arquivo RECEBENDO (.tsv)",
                dataframe_tsv(receiving_rows, RECEIVING_HEADERS),
                file_name="base_recebendo.tsv", mime="text/tab-separated-values",
                disabled=not receiving_rows,
            )
        with right:
            st.download_button(
                "Baixar arquivo EXPEDINDO (.tsv)",
                dataframe_tsv(shipping_rows, SHIPPING_HEADERS),
                file_name="base_expedindo.tsv", mime="text/tab-separated-values",
                disabled=not shipping_rows,
            )
    else:
        st.warning("Nenhum manifesto pôde ser classificado para as duas bases.")
    if messages:
        st.subheader("Avisos para revisar")
        for message in messages:
            st.warning(message)

st.caption(
    "Os endereços sem correspondência ficam em branco e aparecem nos avisos. "
    "Revise os resultados antes de copiar para a planilha online."
)

